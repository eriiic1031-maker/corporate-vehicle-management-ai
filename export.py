"""내보내기: 국세청 업무용승용차 운행기록부(엑셀/PDF), 데이터 백업용 엑셀.

엑셀은 저장소가 아니라 결과물이다. 매번 DB에서 새로 만들기 때문에
기존 운행일지.xlsx의 행 어긋남·병합셀 문제가 생기지 않는다.
"""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import data
import services

WEEK = "월화수목금토일"
HEADERS = ["③사용일자(요일)", "④사용자 부서", "④사용자 성명", "⑤주행 전 계기판의 거리(㎞)", "⑥주행 후 계기판의 거리(㎞)",
           "⑦주행거리(㎞)", "⑧출·퇴근용(㎞)", "⑨일반 업무용(㎞)", "업무용 사용거리(㎞)\n(⑧+⑨)", "⑩비고"]

_thin = Side(style="thin", color="9FB4C7")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _use_date(r: dict) -> date:
    return data.to_date(r.get("use_date")) or data.to_date(r["start_dt"])


def _row_values(r: dict) -> list:
    d = _use_date(r)
    distance = int(r.get("distance") or 0)
    commute = int(r.get("commute_km") or 0)
    personal = int(r.get("personal_km") or 0)
    general = distance - commute - personal
    note = f"예약#{r['id']} {r.get('destination') or ''} {r.get('purpose') or ''}".strip()
    if personal:
        note += f" / 개인용도 {personal:,}km"
    return [f"{d:%Y-%m-%d}({WEEK[d.weekday()]})", r.get("dept") or "", r.get("user_name") or "",
            int(r.get("checkout_km") or 0), int(r.get("return_km") or 0), distance, commute, general, commute + general, note]


def _vehicle_sheet(wb: Workbook, v: dict, rows: list[dict], start: date, end: date, settings: dict):
    ws = wb.create_sheet(services_sheet_name(v["vehicle_no"]))
    bold = Font(bold=True)
    ws["A1"] = "【업무용승용차 운행기록부】"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:J1")
    ws["A1"].alignment = CENTER
    ws["A2"], ws["B2"] = "과세기간", f"{start:%Y.%m.%d} ~ {end:%Y.%m.%d}"
    ws["E2"], ws["F2"] = "상호명", settings.get("company_name", "")
    ws["A3"], ws["B3"] = "사업자등록번호", settings.get("business_no", "")
    ws["A4"] = "1. 기본정보"
    ws["A4"].font = bold
    ws["A5"], ws["B5"] = "①차종", v["model"]
    ws["D5"], ws["E5"] = "②자동차등록번호", v["vehicle_no"]
    ws["G5"], ws["H5"] = "구분", services.vehicle_use_label(v["kind"])
    ws["I5"], ws["J5"] = "현장", v["site"]
    ws["A7"] = "2. 업무용 사용비율 계산"
    ws["A7"].font = bold
    for c in ("A2", "A3", "E2", "A5", "D5", "G5", "I5"):
        ws[c].fill = LABEL_FILL
        ws[c].font = bold
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(8, i, h)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, bold, CENTER, BORDER
    r_idx = 9
    total = business = 0
    monthly = {m: [0, 0] for m in range(1, 13)}
    for r in rows:
        vals = _row_values(r)
        for i, val in enumerate(vals, 1):
            cell = ws.cell(r_idx, i, val)
            cell.border = BORDER
            if 4 <= i <= 9:
                cell.number_format = "#,##0"
        total += vals[5]
        business += vals[8]
        m = _use_date(r).month
        monthly[m][0] += vals[5]
        monthly[m][1] += vals[8]
        r_idx += 1
    if not rows:
        ws.cell(r_idx, 1, "해당 기간 운행 기록 없음")
        r_idx += 1
    t = r_idx + 1
    ws.cell(t, 1, "⑪과세기간 총주행거리(㎞)").fill = LABEL_FILL
    ws.cell(t, 3, total).number_format = "#,##0"
    ws.cell(t, 4, "⑫과세기간 업무용 사용거리(㎞)").fill = LABEL_FILL
    ws.cell(t, 6, business).number_format = "#,##0"
    ws.cell(t, 7, "⑬업무사용비율(⑫/⑪)").fill = LABEL_FILL
    ratio = ws.cell(t, 9, (business / total) if total else 0)
    ratio.number_format = "0.0%"
    for c in (1, 4, 7):
        ws.cell(t, c).font = bold
    # 월별 요약 (오른쪽)
    ws["L4"] = "월별 KM 요약"
    ws["L4"].font = bold
    for i, h in enumerate(["월", "총주행(km)", "업무용(km)", "업무사용비율"], 12):
        cell = ws.cell(5, i, h)
        cell.fill, cell.font, cell.border, cell.alignment = HEAD_FILL, bold, BORDER, CENTER
    for m in range(1, 13):
        row = 5 + m
        tot, bus = monthly[m]
        for i, val in enumerate([f"{m:02d}월", tot, bus, (bus / tot) if tot else 0], 12):
            cell = ws.cell(row, i, val)
            cell.border = BORDER
        ws.cell(row, 13).number_format = ws.cell(row, 14).number_format = "#,##0"
        ws.cell(row, 15).number_format = "0.0%"
    for i, val in enumerate(["합계", total, business, (business / total) if total else 0], 12):
        cell = ws.cell(18, i, val)
        cell.font, cell.border = bold, BORDER
    ws.cell(18, 15).number_format = "0.0%"
    widths = [17, 14, 11, 15, 15, 12, 12, 12, 14, 40, 2, 9, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[8].height = 42
    ws.freeze_panes = "A9"
    return total, business


def services_sheet_name(vno: str) -> str:
    return "".join(ch if ch not in r':\/?*[]' else "_" for ch in vno)[:31] or "차량"


def runlog_excel(vehicle_nos: list[str], start: date, end: date) -> bytes:
    settings = data.get_settings()
    wb = Workbook()
    summary = wb.active
    summary.title = "요약"
    heads = ["차량번호", "차종", "구분", "현장", "운행 건수", "총주행거리(km)", "업무용 사용거리(km)", "업무사용비율"]
    for i, h in enumerate(heads, 1):
        c = summary.cell(1, i, h)
        c.fill, c.font, c.border = HEAD_FILL, Font(bold=True), BORDER
    row = 2
    for vno in vehicle_nos:
        v = services.get_vehicle(vno)
        if not v:
            continue
        rows = services.runlog_rows(vno, start, end)
        total, business = _vehicle_sheet(wb, v, rows, start, end, settings)
        vals = [vno, v["model"], services.vehicle_use_label(v["kind"]), v["site"], len(rows), total, business,
                (business / total) if total else 0]
        for i, val in enumerate(vals, 1):
            summary.cell(row, i, val).border = BORDER
        summary.cell(row, 6).number_format = summary.cell(row, 7).number_format = "#,##0"
        summary.cell(row, 8).number_format = "0.0%"
        row += 1
    for i, w in enumerate([12, 12, 12, 8, 10, 15, 18, 12], 1):
        summary.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _korean_font() -> str | None:
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return None
    for path in (Path("C:/Windows/Fonts/malgun.ttf"), Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
                 Path("/Library/Fonts/AppleGothic.ttf")):
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont("KR", str(path)))
                return "KR"
            except Exception:
                continue
    return None


def runlog_pdf(vehicle_nos: list[str], start: date, end: date) -> bytes | None:
    """PDF 운행기록부. reportlab 또는 한글 글꼴이 없으면 None."""
    font = _korean_font()
    if not font:
        return None
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    settings = data.get_settings()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm)
    title = ParagraphStyle("t", fontName=font, fontSize=15, leading=20, alignment=1)
    normal = ParagraphStyle("n", fontName=font, fontSize=9, leading=12)
    cell = ParagraphStyle("c", fontName=font, fontSize=7.5, leading=9)
    story = []
    for n, vno in enumerate(vehicle_nos):
        v = services.get_vehicle(vno)
        if not v:
            continue
        rows = services.runlog_rows(vno, start, end)
        if n:
            story.append(PageBreak())
        story += [Paragraph("업무용승용차 운행기록부", title), Spacer(1, 4 * mm),
                  Paragraph(f"과세기간: {start:%Y.%m.%d} ~ {end:%Y.%m.%d} &nbsp;&nbsp; 상호명: {settings.get('company_name', '')}"
                            f" &nbsp;&nbsp; 사업자등록번호: {settings.get('business_no', '')}", normal),
                  Paragraph(f"①차종: {v['model']} &nbsp;&nbsp; ②자동차등록번호: {vno} &nbsp;&nbsp; "
                            f"구분: {services.vehicle_use_label(v['kind'])} &nbsp;&nbsp; 현장: {v['site']}", normal),
                  Spacer(1, 3 * mm)]
        table = [[Paragraph(h.replace("\n", "<br/>"), cell) for h in HEADERS]]
        total = business = 0
        for r in rows:
            vals = _row_values(r)
            total += vals[5]
            business += vals[8]
            table.append([Paragraph(f"{x:,}" if isinstance(x, int) else str(x), cell) for x in vals])
        if not rows:
            table.append([Paragraph("해당 기간 운행 기록 없음", cell)] + [""] * 9)
        ratio = f"{business / total:.1%}" if total else "0.0%"
        table.append([Paragraph("⑪총주행거리", cell), "", "", "", "", Paragraph(f"{total:,}", cell),
                      Paragraph("⑫업무용", cell), "", Paragraph(f"{business:,}", cell), Paragraph(f"⑬업무사용비율 {ratio}", cell)])
        t = Table(table, colWidths=[27 * mm, 25 * mm, 20 * mm, 24 * mm, 24 * mm, 20 * mm, 20 * mm, 20 * mm, 24 * mm, 73 * mm],
                  repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9FB4C7")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F2F2F2")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(t)
    if not story:
        story.append(Paragraph("선택한 차량이 없습니다.", normal))
    doc.build(story)
    return buf.getvalue()


def data_backup_excel(emp: dict) -> bytes:
    """권한 범위 안의 원본 데이터를 시트별로 내보낸다."""
    sites = services.manageable_sites(emp)
    ph = ",".join("?" * len(sites))
    frames = {
        "차량": data.query_df(f"SELECT * FROM vehicles WHERE site IN ({ph})", sites),
        "예약": data.query_df(f"SELECT r.* FROM reservations r JOIN vehicles v USING(vehicle_no) WHERE v.site IN ({ph}) ORDER BY r.id", sites),
        "정비": data.query_df(f"SELECT m.* FROM maintenance m JOIN vehicles v USING(vehicle_no) WHERE v.site IN ({ph})", sites),
        "사고·손상": data.query_df(f"SELECT i.* FROM incidents i JOIN vehicles v USING(vehicle_no) WHERE v.site IN ({ph})", sites),
        "비용": data.query_df(f"SELECT e.* FROM expenses e JOIN vehicles v USING(vehicle_no) WHERE v.site IN ({ph})", sites),
    }
    if services.is_admin(emp):
        frames["직원"] = data.query_df("SELECT * FROM employees")
        frames["감사로그"] = data.query_df("SELECT * FROM audit_log ORDER BY id DESC")
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in frames.items():
            df.to_excel(w, sheet_name=name, index=False)
    return buf.getvalue()
