"""업무 규칙 테스트: 상태 흐름, 시간 단위 충돌, 권한, 자동 승인, 정비 일정, 반납/운행기록부."""
import io
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook

import ai
import data
import export
import notify
import services
from services import ServiceError


def emp(name):
    return data.query_one("SELECT * FROM employees WHERE name=?", (name,))


def next_bday(after_days=1):
    d = date.today() + timedelta(days=after_days)
    while not services.is_business_day(d):
        d += timedelta(days=1)
    return d


def at(d, hh, mm=0):
    return datetime.combine(d, time(hh, mm))


@pytest.fixture()
def car(fresh_db):
    admin = emp("김재열")
    services.save_vehicle(admin, {"vehicle_no": "00테9999", "model": "카니발", "kind": "렌트", "site": "평택", "seats": 9,
                                  "fuel_type": "디젤", "body_type": "승합", "odometer": 50000, "rent_start": "2025-01-01",
                                  "rent_end": "2028-12-31", "last_oil_km": 40500, "last_tire_km": 20000,
                                  "key_location": "평택 총무 데스크"}, is_new=True)
    return "00테9999"


def test_migration_and_lifecycle(fresh_db):
    assert data.query_one("SELECT COUNT(*) n FROM vehicles")["n"] == 48
    services.run_scheduled_jobs(force=True)
    counts = {r["status"]: r["n"] for r in data.query_all("SELECT status, COUNT(*) n FROM reservations GROUP BY status")}
    # 원본 데이터: 시작일 지난 승인대기 3건 → 만료, 종료일 지난 승인완료(미반납) 15건 → 연체
    assert counts.get(data.ST_EXPIRED) == 3
    assert counts.get(data.ST_OVERDUE) == 15
    # 정비 시작일 없던 정비중 차량은 이관일로 시작일이 채워진다
    assert data.query_one("SELECT COUNT(*) n FROM maintenance WHERE start_date IS NULL")["n"] == 0


def test_hourly_conflict(car):
    d = next_bday(2)
    a, b, c = emp("김민수"), emp("이지훈"), emp("신현우")
    # 원본 데이터의 진행 중 예약이 정책 검사에 걸리지 않도록 정리
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no IN (?,?,?)", (a["emp_no"], b["emp_no"], c["emp_no"]))
    services.create_reservation(a, car, at(d, 9), at(d, 12), purpose="회의", destination="평택", passengers=2)
    # 같은 날 오후는 다른 사람이 예약 가능 (시간 단위 충돌 검사)
    services.create_reservation(b, car, at(d, 12), at(d, 18), purpose="회의", destination="평택", passengers=2)
    with pytest.raises(ServiceError, match="겹칩니다"):
        services.create_reservation(c, car, at(d, 11), at(d, 13), purpose="x", destination="y", passengers=1)


def test_user_policy_relaxed(car, fresh_db):
    u = emp("김민수")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d1, d2 = next_bday(2), next_bday(5)
    services.create_reservation(u, car, at(d1, 9), at(d1, 12), purpose="a", destination="b", passengers=1)
    # 미래 예약 여러 건 가능 (기존 1인 1예약 완화)
    rid2, _ = services.create_reservation(u, car, at(d2, 9), at(d2, 12), purpose="a", destination="b", passengers=1)
    assert rid2
    # 같은 시간대 본인 중복은 불가
    other = services.recommend("평택", at(d1, 10), at(d1, 11), 1)
    with pytest.raises(ServiceError, match="본인 예약"):
        services.create_reservation(u, other[0].vehicle["vehicle_no"], at(d1, 10), at(d1, 11), purpose="a", destination="b", passengers=1)


def test_permissions(car):
    u, pyeongtaek, bonsa = emp("김민수"), emp("평택현장담당"), emp("본사현장담당")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d = next_bday(3)
    rid, _ = services.create_reservation(u, car, at(d, 9), at(d, 18), purpose="a", destination="b", passengers=1)
    with pytest.raises(ServiceError):
        services.approve(u, rid)          # 일반 사용자 승인 불가
    with pytest.raises(ServiceError):
        services.approve(bonsa, rid)      # 다른 현장 담당자 승인 불가
    services.approve(pyeongtaek, rid)     # 해당 현장 담당자는 승인 가능
    assert services.get_reservation(rid)["status"] == data.ST_APPROVED
    with pytest.raises(ServiceError):
        services.save_settings(pyeongtaek, {"auto_approve_enabled": True})  # 설정은 총무 관리자만
    with pytest.raises(ServiceError):
        services.delete_vehicle(pyeongtaek, car)


def test_reject_needs_reason_and_notifies(car):
    u, admin = emp("김민수"), emp("김재열")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d = next_bday(3)
    rid, _ = services.create_reservation(u, car, at(d, 9), at(d, 18), purpose="a", destination="b", passengers=1)
    with pytest.raises(ServiceError, match="사유"):
        services.reject(admin, rid, "  ")
    services.reject(admin, rid, "임원 행사 차량 회수")
    r = services.get_reservation(rid)
    assert r["status"] == data.ST_REJECTED and r["reject_reason"] == "임원 행사 차량 회수"
    assert any("임원 행사" in n["body"] for n in notify.inbox(u))


def test_auto_approve(car):
    u, admin = emp("김민수"), emp("김재열")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    services.save_settings(admin, {"auto_approve_enabled": True, "auto_approve_time": "17:00"})
    d = next_bday(4)
    rid, msg = services.create_reservation(u, car, at(d, 9), at(d, 18), purpose="a", destination="b", passengers=1)
    assert services.get_reservation(rid)["status"] == data.ST_PENDING  # 아직 전날 17시 전
    before_cutoff = datetime.combine(d - timedelta(days=1), time(16, 30))
    services.run_scheduled_jobs(force=True, now=before_cutoff)
    assert services.get_reservation(rid)["status"] == data.ST_PENDING
    after_cutoff = datetime.combine(d - timedelta(days=1), time(17, 0))
    services.run_scheduled_jobs(force=True, now=after_cutoff)
    r = services.get_reservation(rid)
    assert r["status"] == data.ST_APPROVED and r["auto_approved"] == 1
    # OFF면 자동 승인하지 않는다
    services.save_settings(admin, {"auto_approve_enabled": False})
    d2 = next_bday(6)
    rid2, _ = services.create_reservation(u, car, at(d2, 9), at(d2, 18), purpose="a", destination="b", passengers=1)
    services.run_scheduled_jobs(force=True, now=datetime.combine(d2 - timedelta(days=1), time(18, 0)))
    assert services.get_reservation(rid2)["status"] == data.ST_PENDING


def test_full_flow_checkout_return(car, monkeypatch):
    u, admin = emp("김민수"), emp("김재열")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d = next_bday(1)
    rid, _ = services.create_reservation(u, car, at(d, 9), at(d, 18), purpose="외근", destination="평택공장", passengers=3)
    services.approve(admin, rid)
    monkeypatch.setattr(data, "now", lambda: at(d, 8, 30))
    services.checkout(u, rid)
    r = services.get_reservation(rid)
    assert r["status"] == data.ST_OUT and r["checkout_km"] == 50000
    monkeypatch.setattr(data, "now", lambda: at(d, 17, 30))  # 17:30에 반납
    with pytest.raises(ServiceError):
        services.return_vehicle(u, rid, return_km=49000, commute_km=0, use_date=d, fuel_level="1/2", parking_location="B2")
    with pytest.raises(ServiceError, match="주차"):
        services.return_vehicle(u, rid, return_km=50120, commute_km=0, use_date=d, fuel_level="1/2", parking_location="")
    msg = services.return_vehicle(
        u, rid, return_km=50120, commute_km=20, personal_km=10, use_date=d, fuel_level="1/2", parking_location="지하2층 B-12",
        odometer_photo=("odo.jpg", b"fake"), damage={"itype": "손상", "description": "뒷범퍼 긁힘", "photos": [("d.jpg", b"x")]},
        expenses=[{"etype": "유류비", "amount": 50000, "receipt": ("r.jpg", b"r")}, {"etype": "주차비", "amount": 0}])
    r = services.get_reservation(rid)
    assert r["status"] == data.ST_DONE
    assert (r["distance"], r["commute_km"], r["personal_km"], r["business_km"]) == (120, 20, 10, 110)
    v = services.get_vehicle(car)
    assert v["odometer"] == 50120 and v["parking_location"] == "지하2층 B-12"
    assert data.query_one("SELECT COUNT(*) n FROM incidents")["n"] == 1
    assert data.query_one("SELECT SUM(amount) s FROM expenses")["s"] == 50000
    assert "업무용 110" in msg

    # 운행기록부 엑셀: 5행 기본정보, 8행 헤더, 9행 데이터, 업무사용비율 계산
    xls = export.runlog_excel([car], d, d)
    wb = load_workbook(io.BytesIO(xls))
    ws = wb[car]
    assert ws["E5"].value == car and ws["B5"].value == "카니발"
    assert ws.cell(9, 6).value == 120 and ws.cell(9, 7).value == 20 and ws.cell(9, 8).value == 90 and ws.cell(9, 9).value == 110
    assert abs(wb["요약"].cell(2, 8).value - 110 / 120) < 1e-9
    pdf = export.runlog_pdf([car], d, d)
    assert pdf is None or pdf[:4] == b"%PDF"

    # 과태료 운전자 조회
    found = services.find_driver(car, at(d, 13))
    assert found and found[0]["user_name"] == "김민수"
    assert not services.find_driver(car, at(d, 18))   # 반납(17:30) 이후는 해당 없음


def test_noshow_and_overdue(car):
    u, admin = emp("김민수"), emp("김재열")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d = next_bday(1)
    rid, _ = services.create_reservation(u, car, at(d, 9), at(d, 18), purpose="a", destination="b", passengers=1)
    services.approve(admin, rid)
    services.run_scheduled_jobs(force=True, now=at(d, 11, 30))   # 시작 + 2시간 경과, 출차 없음
    assert services.get_reservation(rid)["status"] == data.ST_NOSHOW

    d2 = next_bday(3)
    rid2, _ = services.create_reservation(u, car, at(d2, 9), at(d2, 18), purpose="a", destination="b", passengers=1)
    services.approve(admin, rid2)
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status=?, checkout_at=?, checkout_km=50000 WHERE id=?",
                     (data.ST_OUT, f"{d2} 09:00", rid2))
    services.run_scheduled_jobs(force=True, now=at(d2, 19))
    assert services.get_reservation(rid2)["status"] == data.ST_OVERDUE
    # 연체 중에는 새 예약 불가
    with pytest.raises(ServiceError, match="연체"):
        services.create_reservation(u, car, at(next_bday(6), 9), at(next_bday(6), 12), purpose="a", destination="b", passengers=1)


def test_cancel_and_change_request(car):
    u, admin = emp("김민수"), emp("김재열")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d = next_bday(3)
    rid, _ = services.create_reservation(u, car, at(d, 9), at(d, 12), purpose="a", destination="b", passengers=1)
    assert "바로" in services.request_change(u, rid, at(d, 13), at(d, 15))   # 승인 전: 즉시 변경
    assert services.get_reservation(rid)["start_dt"].endswith("13:00")
    services.approve(admin, rid)
    assert "요청" in services.request_change(u, rid, at(d, 14), at(d, 17), "회의 연장")
    cr = data.query_one("SELECT * FROM change_requests WHERE reservation_id=?", (rid,))
    services.decide_change(admin, cr["id"], True)
    assert services.get_reservation(rid)["end_dt"].endswith("17:00")
    services.cancel(u, rid, "일정 취소")
    assert services.get_reservation(rid)["status"] == data.ST_CANCELED


def test_maintenance_schedule_blocks_period(car):
    admin, u = emp("김재열"), emp("김민수")
    with data.tx() as conn:
        conn.execute("UPDATE reservations SET status='반납완료' WHERE emp_no=?", (u["emp_no"],))
    d = next_bday(3)
    services.add_maintenance(admin, car, d, d, "엔진오일")
    with pytest.raises(ServiceError, match="정비"):
        services.create_reservation(u, car, at(d, 9), at(d, 12), purpose="a", destination="b", passengers=1)
    d2 = next_bday((d - date.today()).days + 1)
    services.create_reservation(u, car, at(d2, 9), at(d2, 12), purpose="a", destination="b", passengers=1)  # 정비 끝난 뒤는 가능


def test_vehicle_alerts(car):
    v = services.get_vehicle(car)
    alerts = services.vehicle_alerts(v, data.get_settings())
    types = {a["type"]: a["level"] for a in alerts}
    assert types.get("엔진오일·필터") == "임박"   # 마지막 교체 후 9,500 km (기준 10,000의 90% 이상)
    assert "타이어" not in types                # 30,000 km < 36,000 km
